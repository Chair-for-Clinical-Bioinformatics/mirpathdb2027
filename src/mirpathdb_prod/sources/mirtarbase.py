"""miRTarBase target-source adapter (fetch/load, real implementation).

miRTarBase 9.0 (Huang et al. 2022) is a **validated** miRNA-target-interaction (MTI) source.
It is a backbone of the user-facing confidence rollup alongside DIANA-TarBase:
its strong/functional evidence (reporter assay / Western blot / qPCR) maps to
**HIGH** confidence, and every other row maps to the broader experimental band
**MEDIUM** (see ``docs/target_evidence.md`` and ``config/target_sources.yaml``).
"Every other row" includes the rows whose ``Support Type`` cell is empty -- a
missing upstream classification, labelled :data:`BLANK_SUPPORT_LABEL` in the
target-pair artefacts. DIANA-TarBase is the parallel validated resource that
ALSO feeds HIGH + MEDIUM.

Per-organism coverage (organism availability matrix — verified):
    hsa ✓ dominant · mmu ✓ strong · rno ✓ (validated, 3rd species)
    cfa / gga thin/tail; evaluated + EXCLUDED (owner-omitted) · mml ✗ (no miRTarBase
    per-species file; macaque's only validated evidence is a single DIANA-TarBase v9
    MTI, marked)

This module does fetch/load ONLY — no normalization (that is
:mod:`mirpathdb_prod.normalize`) and no pass/fail judgment (that is
:mod:`mirpathdb_prod.validation`). It must never import prototype code.

The per-species file is the official miRTarBase MTI workbook
(``hsa_MTI.xlsx`` / ``mmu_MTI.xlsx``). Each row is one supported interaction with
columns::

    miRTarBase ID | miRNA | Species (miRNA) | Target Gene |
    Target Gene (Entrez ID) | Species (Target Gene) | Experiments |
    Support Type | References (PMID)

The ``miRNA`` column is already in display-name form (``hsa-let-7a-5p``) so it
joins directly to the v2 p-value-matrix row labels. ``Target Gene`` is a gene
SYMBOL (the v2 GeneTrail2 ORA space), and ``Target Gene (Entrez ID)`` carries the
Entrez fallback ID.

Excel autocorrupts some ``Target Gene`` symbol cells into date objects
(``SEPT2`` -> ``datetime(2022, 9, 2)``). :func:`load_mti` repairs those via
:mod:`mirpathdb_prod.sources.gene_identity` (the W1 guard, mirroring the
webserver importer): recover through the workbook's own Entrez column when it
resolves unambiguously, else emit the ``ENTREZ:<id>`` sentinel, else skip and
count -- a stringified timestamp can never leave the loader as a gene symbol.
"""

from __future__ import annotations

import datetime
import warnings
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from mirpathdb_prod.sources import validate_species_column
from mirpathdb_prod.sources.gene_identity import (
    entrez_sentinel,
    is_date_artifact,
    is_valid_gene_symbol,
    normalize_entrez_id,
)

# Source identifier stamped on loaded records (machine-readable mirror lives in
# config/target_sources.yaml under the matching `id`).
SOURCE = "mirtarbase"

# Pinned snapshot the production build targets (see provenance/manifest.yaml).
SOURCE_VERSION = "9.0"

# Granular internal provenance tiers this source can emit (a subset of the
# 5-tier enum in docs/target_evidence.md). miRTarBase carries validated evidence
# only: strong/functional -> val_strong (HIGH); every other row (weak /
# non-functional labels and blank Support Type cells) -> val_weak (MEDIUM).
PROVENANCE_TIERS: tuple[str, ...] = ("val_strong", "val_weak")

# Production organisms with usable miRTarBase per-species MTI coverage
# (canonical taxids). cfa/gga have only a thin tail and were evaluated +
# excluded; macaque (mml 9544) has no miRTarBase per-species file (its only
# validated evidence is a single DIANA-TarBase v9 MTI, marked).
SUPPORTED_TAXIDS: frozenset[int] = frozenset({9606, 10090, 10116})

# Canonical taxid -> the worksheet name miRTarBase uses inside the per-species
# workbook (the species' scientific name).
_TAXID_TO_SHEET: dict[int, str] = {
    9606: "Homo sapiens",
    10090: "Mus musculus",
    10116: "Rattus norvegicus",
}

# Canonical taxid -> short organism code used in the W1 gene-identity summary.
_TAXID_TO_CODE: dict[int, str] = {9606: "hsa", 10090: "mmu", 10116: "rno"}

# Cap on the example skipped-row lines printed by the W1 summary (mirrors the
# webserver importer's MAX_SKIP_EXAMPLES).
MAX_SKIP_EXAMPLES = 10

# Expected column header of the per-species MTI workbook. Verified against the
# cached hsa/mmu files; a reshaped file fails loudly rather than silently
# mis-indexing.
MTI_COLUMNS: tuple[str, ...] = (
    "miRTarBase ID",
    "miRNA",
    "Species (miRNA)",
    "Target Gene",
    "Target Gene (Entrez ID)",
    "Species (Target Gene)",
    "Experiments",
    "Support Type",
    "References (PMID)",
)

# The single Support Type value miRTarBase uses for STRONG functional evidence
# (reporter assay / Western blot / qPCR). Every other row -- the weak and
# non-functional labels ("Functional MTI (Weak)", "Non-Functional MTI", ...) and
# a row with an empty cell (:data:`BLANK_SUPPORT_LABEL`) -- goes to MEDIUM.
STRONG_SUPPORT_TYPE = "Functional MTI"

# The closed Support Type vocabulary miRTarBase 9.0 uses (verified against the
# cached hsa/mmu/rno workbooks). The strong/weak rollup keys off
# :data:`STRONG_SUPPORT_TYPE`; this is the EXPECTED set used by
# :func:`validate_vocabulary` to catch a silent schema/value rename (M1). If
# miRTarBase ships a NEW value it is warned (additive), not silently bucketed
# into "weak" -- but a rename of the strong label itself would zero out HIGH, so
# that case raises. The empty cell is deliberately NOT a member: it is not a
# miRTarBase category (see :data:`BLANK_SUPPORT_LABEL`).
EXPECTED_SUPPORT_TYPES: frozenset[str] = frozenset(
    {
        "Functional MTI",          # STRONG (reporter/Western/qPCR)
        "Functional MTI (Weak)",   # weak experimental
        "Non-Functional MTI",      # reported non-functional
        "Non-Functional MTI (Weak)",
    }
)

# The fifth Support Type state: an EMPTY cell (openpyxl yields None). It lies
# outside the closed vocabulary above, yet it is the majority state of the pinned
# 9.0 workbooks (hsa 454,388 of 957,040 rows = 47.5 %; mmu 57.3 %; rno 54.3 %).
# It is a MISSING upstream classification: the source did not call these rows
# weak, it gave them no classification at all. miRPathDB's evidence-organization
# rule places them in the broader experimental band (MEDIUM) because they are not
# the exact strong label; MEDIUM does not assert functional repression. The
# target-pair exporter writes this text as the miRTarBase provenance detail of
# such rows, so it must stay free of the annotation separators (tab, newline,
# ';', '|').
BLANK_SUPPORT_LABEL = "support classification not provided in the source"

# How blank-support rows are handled; quoted by the loader warning and recorded
# per organism in the target-pair counts.json.
BLANK_SUPPORT_POLICY = "kept in the broader experimental band (MEDIUM)"


def support_type_census(rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    """Count the ``Support Type`` states of loaded miRTarBase rows (no side effects).

    A cell is blank when it is ``None``, empty or whitespace-only; non-empty cells
    are compared after ``strip`` (as in the band builder). Shared by
    :func:`validate_vocabulary` and the target-pair exporter's counts.json census.

    Returns:
        ``n_rows``, ``n_blank``, ``n_nonempty``, ``n_strong`` (rows carrying
        :data:`STRONG_SUPPORT_TYPE`), ``blank_fraction`` (``n_blank / n_rows``,
        0.0 for no rows), ``values`` (sorted distinct non-empty values) and
        ``unknown`` (the sorted subset outside :data:`EXPECTED_SUPPORT_TYPES`).
    """
    n_rows = 0
    n_blank = 0
    n_strong = 0
    seen: set[str] = set()
    for row in rows:
        n_rows += 1
        val = row.get("Support Type")
        text = "" if val is None else str(val).strip()
        if not text:
            n_blank += 1
            continue
        seen.add(text)
        if text == STRONG_SUPPORT_TYPE:
            n_strong += 1
    return {
        "n_rows": n_rows,
        "n_blank": n_blank,
        "n_nonempty": n_rows - n_blank,
        "n_strong": n_strong,
        "blank_fraction": n_blank / n_rows if n_rows else 0.0,
        "values": sorted(seen),
        "unknown": sorted(seen - EXPECTED_SUPPORT_TYPES),
    }


def validate_vocabulary(rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    """Count blank Support Types, warn on unknown values, raise if STRONG vanished (M1).

    Additive guard that does NOT change the rollup. It tallies the loaded rows'
    ``Support Type`` column via :func:`support_type_census` and

    * ``warnings.warn``s ONCE per call when any cell is blank (``None``, empty or
      whitespace), stating the count, the total, the fraction and the policy
      (:data:`BLANK_SUPPORT_POLICY`, provenance label
      :data:`BLANK_SUPPORT_LABEL`). Blank cells never raise;
    * ``warnings.warn``s once per unknown non-empty value (so a renamed/added
      miRTarBase category surfaces instead of being silently bucketed weak);
    * raises ``ValueError`` only when ZERO rows carry the
      :data:`STRONG_SUPPORT_TYPE` label among non-empty support values -- the
      signal that the strong/HIGH evidence label itself was renamed (which would
      silently empty the HIGH tier). A blank-only input has no non-empty value
      and therefore does not raise.

    Returns:
        The :func:`support_type_census` dict.
    """
    census = support_type_census(rows)
    if census["n_blank"]:
        warnings.warn(
            f"miRTarBase: {census['n_blank']:,} of {census['n_rows']:,} rows "
            f"({census['blank_fraction']:.1%}) have an empty Support Type -- a missing "
            f"upstream classification outside EXPECTED_SUPPORT_TYPES; "
            f"{BLANK_SUPPORT_POLICY}; provenance label {BLANK_SUPPORT_LABEL!r}.",
            stacklevel=2,
        )
    for value in census["unknown"]:
        warnings.warn(
            f"miRTarBase: unknown Support Type {value!r} not in "
            f"EXPECTED_SUPPORT_TYPES {sorted(EXPECTED_SUPPORT_TYPES)} -- "
            f"treated as weak by the rollup; confirm the schema did not drift.",
            stacklevel=2,
        )
    if census["n_nonempty"] and census["n_strong"] == 0:
        raise ValueError(
            f"miRTarBase: 0 rows carry the strong label {STRONG_SUPPORT_TYPE!r} "
            f"among {census['n_nonempty']} rows with a Support Type "
            f"(seen: {census['values']}). "
            f"The strong/HIGH evidence label appears to have been renamed -- "
            f"refusing to silently empty the HIGH tier."
        )
    return census


def load_mti(
    path: Path, *, taxid: int, ledger: list[dict[str, Any]] | None = None
) -> list[dict[str, Any]]:
    """Load raw miRNA-target interactions for one organism.

    Near-verbatim load: returns one dict per worksheet row, keyed by the official
    column names in :data:`MTI_COLUMNS`. The ONE exception to "verbatim" is the
    W1 gene-identity guard (:func:`_recover_gene_identity`): ``Target Gene``
    cells that Excel autocorrected into dates (or other non-symbol artifacts)
    are recovered via the workbook's own ``Target Gene (Entrez ID)`` column or
    dropped and counted -- a stringified timestamp can never leave this loader.
    The worksheet is selected by ``taxid`` (its scientific name);
    if the workbook has a single sheet that sheet is used regardless of name.

    Args:
        path: Path to the per-species miRTarBase MTI workbook
            (``*_MTI.xlsx``).
        taxid: Canonical NCBI taxid of the organism to load
            (must be in :data:`SUPPORTED_TAXIDS`).
        ledger: Optional list the W1 guard appends one structured record to per
            AFFECTED row (recovered, sentinel'd, or dropped) -- the loader lane
            of the release exclusion ledger. Keys follow
            :data:`mirpathdb_prod.sources.gene_identity.EXCLUSION_LEDGER_COLUMNS`.
            Clean rows are never recorded. ``None`` (default) disables recording;
            the returned rows are identical either way.

    Returns:
        Raw MTI row dicts (verbatim except the W1 ``Target Gene`` identity
        guard; rows with no recoverable gene identity are dropped and counted).

    Raises:
        ValueError: if ``taxid`` is not supported, the file header does not
            match :data:`MTI_COLUMNS`, or the expected worksheet is absent.
        FileNotFoundError: if ``path`` does not exist.
    """
    import openpyxl  # local import: heavy dep, keeps module import cheap

    if taxid not in SUPPORTED_TAXIDS:
        raise ValueError(
            f"miRTarBase has no per-species MTI file for taxid {taxid}; "
            f"supported: {sorted(SUPPORTED_TAXIDS)}"
        )

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"miRTarBase MTI workbook not found: {path}")

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = _select_sheet(workbook.sheetnames, taxid)
        worksheet = workbook[sheet_name]
        row_iter = worksheet.iter_rows(values_only=True)

        try:
            header = next(row_iter)
        except StopIteration as exc:  # empty worksheet
            raise ValueError(f"{path}: worksheet {sheet_name!r} is empty") from exc

        header_tuple = tuple("" if h is None else str(h) for h in header)
        if header_tuple != MTI_COLUMNS:
            raise ValueError(
                f"{path}: unexpected miRTarBase MTI header on sheet {sheet_name!r}\n"
                f"  expected: {MTI_COLUMNS}\n"
                f"  got:      {header_tuple}"
            )

        rows: list[dict[str, Any]] = []
        for values in row_iter:
            if values is None or all(v is None for v in values):
                continue  # skip fully-blank trailing rows
            rows.append(dict(zip(MTI_COLUMNS, values, strict=False)))
        # Row-level species check: the workbook's own annotations must agree
        # with the requested taxid (protection beyond the sheet-name
        # convention; a single-sheet workbook is accepted regardless of name).
        validate_species_column(
            rows,
            column="Species (miRNA)",
            taxid=taxid,
            path=path,
            source_label="miRTarBase",
        )
        # W1: gene-symbol identity guard + Entrez recovery on the Target Gene
        # column (Excel date-autocorrect corruption; 2026-08-24 audit defect (b)).
        rows = _recover_gene_identity(
            rows, species=_TAXID_TO_CODE.get(taxid, str(taxid)), ledger=ledger
        )
        # M1: additive file-format vocab check (warns once with the blank
        # Support Type count, warns on unknown values, raises only if the STRONG
        # label vanished). Does NOT alter the rows: a blank cell stays None.
        validate_vocabulary(rows)
        return rows
    finally:
        workbook.close()


def _select_sheet(sheet_names: list[str], taxid: int) -> str:
    """Pick the worksheet for ``taxid`` (or the sole sheet if there is only one)."""
    if len(sheet_names) == 1:
        return sheet_names[0]
    wanted = _TAXID_TO_SHEET.get(taxid)
    if wanted is not None and wanted in sheet_names:
        return wanted
    raise ValueError(
        f"no worksheet for taxid {taxid} (wanted {wanted!r}); "
        f"available sheets: {sheet_names}"
    )


def _recover_gene_identity(
    rows: list[dict[str, Any]],
    *,
    species: str,
    ledger: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Enforce gene-symbol identity on the ``Target Gene`` column (W1 guard).

    Excel autocorrects symbol cells into date objects (``SEPT2`` ->
    ``datetime(2022, 9, 2)``); read back verbatim they ship stringified
    timestamps as gene symbols. Mirroring the webserver importer
    (``mirpathdb.management.commands.import_targets``):

    * a cell that passes :func:`gene_identity.is_valid_gene_symbol` is kept
      verbatim (and its Entrez ID recorded in the in-workbook resolution map);
    * an invalid cell (date/datetime object, stringified date/timestamp,
      numeric, placeholder) is recovered via the row's own
      ``Target Gene (Entrez ID)``: rewritten to the symbol this same workbook
      uses for that Entrez ID when UNAMBIGUOUS (exactly one valid symbol), else
      to the ``ENTREZ:<id>`` sentinel;
    * a row with neither a valid symbol nor a usable Entrez ID is dropped and
      counted -- never silently, and never as a stringified timestamp.

    Prints one per-species summary line (plus skip reasons/examples when any
    row was dropped), mirroring the importer's accounting. When ``ledger`` is a
    list, one structured record (keys =
    :data:`gene_identity.EXCLUSION_LEDGER_COLUMNS`) is appended per AFFECTED row:
    disposition ``recovered_symbol`` / ``recovered_sentinel`` for repaired rows,
    ``excluded`` for dropped ones. Clean rows are never recorded.
    """
    rows_read = len(rows)
    date_cells = 0            # Target Gene cells that are datetime/date objects
    invalid_cells = 0         # non-empty, non-date cells failing is_valid_gene_symbol
    skipped: Counter[str] = Counter()
    skip_examples: list[str] = []
    # invalid rows deferred for post-pass Entrez recovery: (row, raw, entrez, is_date)
    pending: list[tuple[dict[str, Any], object, str, bool]] = []
    # entrez id -> valid symbols this workbook uses for it (resolution map)
    ent_symbols: dict[str, set[str]] = {}
    out: list[dict[str, Any]] = []

    def _record(
        row: dict[str, Any],
        *,
        mirna: str,
        raw: object,
        entrez: str | None,
        disposition: str,
        reason: str,
        recovered: str = "",
        syms: set[str] | frozenset[str] = frozenset(),
    ) -> None:
        """Append one loader-lane ledger record (no-op when ledger is None)."""
        if ledger is None:
            return
        ledger.append({
            "record_id": f"{species}-w1-{len(ledger) + 1:05d}",
            "organism": species,
            "mirna": mirna,
            "raw_value": "" if raw is None else str(raw).strip(),
            "entrez": entrez or "",
            "multiplicity": 1,
            "n_source_rows_without_entrez": 0 if entrez else 1,
            "annotation_symbols": "|".join(sorted(syms)),
            "n_annotation_symbols": len(syms),
            "recovered_symbol": recovered,
            "atomic_evidence": "",  # band unknown at load time (export lane fills it)
            "date_form_check": "date" if is_date_artifact(raw) else "non_date",
            "disposition": disposition,
            "reason": reason,
            "support_types": str(row.get("Support Type") or "").strip(),
            "workbook_rows": 1,
            "mirtarbase_ids": str(row.get("miRTarBase ID") or "").strip(),
        })

    def _skip(reason: str, mirna: str, raw: object, row: dict[str, Any]) -> None:
        skipped[reason] += 1
        if len(skip_examples) < MAX_SKIP_EXAMPLES:
            skip_examples.append(
                f"miRNA={mirna or '<empty>'} gene={raw!r} reason={reason}"
            )
        _record(
            row,
            mirna=mirna,
            raw=raw,
            entrez=normalize_entrez_id(row.get("Target Gene (Entrez ID)")),
            disposition="excluded",
            reason=reason,
        )

    for row in rows:
        raw_mirna = row.get("miRNA")
        mirna = str(raw_mirna).strip() if raw_mirna is not None else ""
        raw_gene = row.get("Target Gene")
        if not mirna:
            _skip("empty_mirna", mirna, raw_gene, row)
            continue
        is_date = isinstance(raw_gene, (datetime.datetime, datetime.date))
        gene = "" if raw_gene is None or is_date else str(raw_gene).strip()
        if not is_date and not gene:
            _skip("empty_gene", mirna, raw_gene, row)
            continue
        entrez = normalize_entrez_id(row.get("Target Gene (Entrez ID)"))
        if is_date or not is_valid_gene_symbol(raw_gene):
            if is_date:
                date_cells += 1
            else:
                invalid_cells += 1
            if entrez:
                out.append(row)  # kept; Target Gene rewritten in the post-pass
                pending.append((row, raw_gene, entrez, is_date))
            else:
                _skip("invalid_gene_no_entrez", mirna, raw_gene, row)
            continue
        if entrez:
            ent_symbols.setdefault(entrez, set()).add(gene)
        out.append(row)

    # Post-pass recovery: the resolution map is only complete after the full
    # sweep, because Excel corrupts EVERY occurrence of an affected symbol --
    # the valid spelling, if any, comes from rows using the modern alias
    # (e.g. SEPTIN2 for 4735).
    recovered_symbol = 0    # resolved to the workbook's own symbol for that ID
    recovered_sentinel = 0  # ENTREZ:<id> (no unambiguous symbol in workbook)
    dates_recovered = 0
    for row, raw_gene, entrez, was_date in pending:
        syms = ent_symbols.get(entrez, set())
        if len(syms) == 1:
            row["Target Gene"] = next(iter(syms))
            recovered_symbol += 1
            disposition = "recovered_symbol"
        else:  # no valid symbol anywhere in the workbook, or several (ambiguous)
            row["Target Gene"] = entrez_sentinel(entrez)
            recovered_sentinel += 1
            disposition = "recovered_sentinel"
        if was_date:
            dates_recovered += 1
        mirna_s = str(row.get("miRNA") or "").strip()
        _record(
            row,
            mirna=mirna_s,
            raw=raw_gene,
            entrez=entrez,
            disposition=disposition,
            reason="date_cell" if was_date else "invalid_cell",
            recovered=row["Target Gene"],
            syms=syms,
        )

    print(
        f"[mirtarbase] {species}: rows read {rows_read:,} | "
        f"date cells {date_cells} (recovered via Entrez {dates_recovered}) | "
        f"other invalid {invalid_cells} | recovered to symbol {recovered_symbol} | "
        f"stored as ENTREZ:<id> {recovered_sentinel} | skipped {sum(skipped.values())}",
        flush=True,
    )
    if skipped:
        print(f"[mirtarbase] {species}: skip reasons {dict(skipped)}", flush=True)
        for ex in skip_examples:
            print(f"  skipped: {ex}", flush=True)
    return out
